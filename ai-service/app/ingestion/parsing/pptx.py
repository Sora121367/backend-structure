from collections.abc import Iterator
from pathlib import Path
from zipfile import BadZipFile

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError
from pptx.shapes.base import BaseShape

from app.ingestion.parsing._text import normalize, table_text
from app.ingestion.parsing.errors import DocumentParseError
from app.ingestion.parsing.models import Anchor, Block, ParsedDocument


def _reading_order(shapes) -> list[BaseShape]:
    """Top to bottom, then left to right; shape order in the file is z-order, not reading order."""
    return sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))


def _walk(shapes) -> Iterator[BaseShape]:
    for shape in _reading_order(shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _slide_blocks(slide, number: int) -> tuple[list[Block], bool]:
    """Blocks of one slide, and whether it holds pictures but no text besides the title."""
    title_shape = slide.shapes.title
    title = normalize(title_shape.text_frame.text) if title_shape is not None else ""
    anchor = Anchor(slide=number, section_path=(title,) if title else ())

    blocks = [Block(kind="heading", text=title, heading_level=1, anchor=anchor)] if title else []
    has_picture = False
    for shape in _walk(slide.shapes):
        if title_shape is not None and shape.shape_id == title_shape.shape_id:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            has_picture = True
        if shape.has_text_frame:
            # One block per text box: a title wrapped over two lines stays one block.
            lines = (normalize(p.text) for p in shape.text_frame.paragraphs)
            if text := "\n".join(line for line in lines if line):
                blocks.append(Block(kind="paragraph", text=text, anchor=anchor))
        elif getattr(shape, "has_table", False) and shape.has_table:
            rows = [[normalize(cell.text) for cell in row.cells] for row in shape.table.rows]
            if text := table_text(rows):
                blocks.append(Block(kind="table", text=text, anchor=anchor))

    needs_ocr = has_picture and len(blocks) == (1 if title else 0)
    if slide.has_notes_slide and (notes := normalize(slide.notes_slide.notes_text_frame.text)):
        blocks.append(Block(kind="notes", text=notes, anchor=anchor))
    return blocks, needs_ocr


def parse_pptx(path: Path) -> ParsedDocument:
    try:
        presentation = pptx.Presentation(str(path))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError) as e:
        raise DocumentParseError("The file is not a readable PowerPoint presentation.") from e

    blocks: list[Block] = []
    needs_ocr: list[int] = []
    for number, slide in enumerate(presentation.slides, start=1):
        slide_blocks, slide_needs_ocr = _slide_blocks(slide, number)
        blocks.extend(slide_blocks)
        if slide_needs_ocr:
            needs_ocr.append(number)

    if not blocks and not needs_ocr:
        raise DocumentParseError("The presentation has no text or pictures.")
    return ParsedDocument(
        source_type="pptx",
        title=normalize(presentation.core_properties.title or "") or None,
        blocks=tuple(blocks),
        slide_count=len(presentation.slides),
        pages_needing_ocr=tuple(needs_ocr),
    )
