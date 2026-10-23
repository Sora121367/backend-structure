"""Images to read for the pages a parser flagged in `pages_needing_ocr`."""

from collections.abc import Iterator
from pathlib import Path

import pptx
import pymupdf
from pptx.enum.shapes import MSO_SHAPE_TYPE

_PDF_DPI = 200


def pdf_page_images(path: Path, page_numbers: tuple[int, ...]) -> Iterator[tuple[int, list[bytes]]]:
    """Each flagged page rendered whole: scans hold the text in the picture, not beside it."""
    with pymupdf.open(path, filetype="pdf") as doc:
        for number in page_numbers:
            pixmap = doc[number - 1].get_pixmap(dpi=_PDF_DPI)
            yield number, [pixmap.tobytes("png")]


def _pictures(shapes) -> Iterator[bytes]:
    for shape in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape.image.blob


def pptx_slide_images(
    path: Path, slide_numbers: tuple[int, ...]
) -> Iterator[tuple[int, list[bytes]]]:
    """The pictures on each flagged slide, in reading order. Slides aren't rendered (no office
    suite on the server), so text drawn as shapes can't be read this way."""
    slides = pptx.Presentation(str(path)).slides
    for number in slide_numbers:
        yield number, list(_pictures(slides[number - 1].shapes))
