import io

import pptx
import pymupdf
import pytest
from pptx.util import Inches

from app.ingestion.parsing import DocumentParseError, parse_document

TITLE_AND_CONTENT, TITLE_ONLY, BLANK = 1, 5, 6


def save(presentation, tmp_path):
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    return path


def png() -> io.BytesIO:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pix.clear_with(200)
    return io.BytesIO(pix.tobytes("png"))


def content_slide(presentation, title: str, body: str):
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_AND_CONTENT])
    slide.shapes.title.text = title
    slide.placeholders[1].text_frame.text = body
    return slide


def test_slides_become_anchored_blocks(tmp_path):
    presentation = pptx.Presentation()
    presentation.core_properties.title = "Week 1"
    first = content_slide(presentation, "Variables", "Names for values\nAssigned with =")
    first.notes_slide.notes_text_frame.text = "Mention naming rules"
    content_slide(presentation, "Loops", "for and while")
    parsed = parse_document(save(presentation, tmp_path), "pptx")

    assert parsed.title == "Week 1"
    assert parsed.slide_count == 2
    assert [(b.anchor.slide, b.kind, b.text, b.anchor.section_path) for b in parsed.blocks] == [
        (1, "heading", "Variables", ("Variables",)),
        (1, "paragraph", "Names for values\nAssigned with =", ("Variables",)),
        (1, "notes", "Mention naming rules", ("Variables",)),
        (2, "heading", "Loops", ("Loops",)),
        (2, "paragraph", "for and while", ("Loops",)),
    ]
    assert parsed.pages_needing_ocr == ()


def test_reads_tables_and_grouped_shapes_in_reading_order(tmp_path):
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for (r, c), text in {
        (0, 0): "Term",
        (0, 1): "Meaning",
        (1, 0): "API",
        (1, 1): "Interface",
    }.items():
        table.cell(r, c).text = text
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Top"
    parsed = parse_document(save(presentation, tmp_path), "pptx")

    assert [(b.kind, b.text) for b in parsed.blocks] == [
        ("paragraph", "Top"),
        ("table", "Term | Meaning\nAPI | Interface"),
    ]
    assert parsed.blocks[0].anchor.section_path == ()


def test_flags_picture_slides_for_ocr(tmp_path):
    presentation = pptx.Presentation()
    content_slide(presentation, "Text slide", "body")
    diagram = presentation.slides.add_slide(presentation.slide_layouts[TITLE_ONLY])
    diagram.shapes.title.text = "Architecture"
    diagram.shapes.add_picture(png(), Inches(1), Inches(2))
    scan = presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    scan.shapes.add_picture(png(), Inches(1), Inches(1))
    parsed = parse_document(save(presentation, tmp_path), "pptx")

    assert parsed.pages_needing_ocr == (2, 3)


def test_rejects_presentation_without_content(tmp_path):
    presentation = pptx.Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[BLANK])
    with pytest.raises(DocumentParseError):
        parse_document(save(presentation, tmp_path), "pptx")


def test_rejects_word_file_declared_as_pptx(tmp_path):
    import docx

    path = tmp_path / "wrong.pptx"
    document = docx.Document()
    document.add_paragraph("hi")
    document.save(str(path))
    with pytest.raises(DocumentParseError):
        parse_document(path, "pptx")


@pytest.mark.parametrize("data", [b"", b"not a pptx"])
def test_rejects_unreadable_files(tmp_path, data):
    path = tmp_path / "bad.pptx"
    path.write_bytes(data)
    with pytest.raises(DocumentParseError):
        parse_document(path, "pptx")
