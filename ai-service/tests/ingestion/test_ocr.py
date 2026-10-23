import importlib.util
import io

import pptx
import pymupdf
import pytest
from pptx.util import Inches

from app.ingestion.ocr import run_ocr
from app.ingestion.parsing import parse_document


class FakeEngine:
    """Returns canned lines per call, in order; remembers the images it was given."""

    def __init__(self, *results: list[str]) -> None:
        self._results = list(results)
        self.images: list[bytes] = []

    def recognize(self, image: bytes) -> list[str]:
        self.images.append(image)
        return self._results.pop(0) if self._results else []


def png_bytes() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pix.clear_with(200)
    return pix.tobytes("png")


def pdf_with_scan(tmp_path, outline=None):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "Text page")
    doc.new_page().insert_image(pymupdf.Rect(72, 72, 272, 272), stream=png_bytes())
    doc.new_page().insert_text((72, 100), "Last page")
    if outline:
        doc.set_toc(outline)
    path = tmp_path / "scan.pdf"
    doc.save(path)
    return path


def test_pdf_scan_page_gets_ocr_text_with_page_anchor(tmp_path):
    path = pdf_with_scan(tmp_path, outline=[[1, "Chapter 1", 1]])
    parsed = parse_document(path, "pdf")
    engine = FakeEngine(["First line", "  Second   line "])

    result = run_ocr(path, parsed, engine)

    assert result.pages_needing_ocr == ()
    assert [(b.anchor.page, b.text) for b in result.blocks] == [
        (1, "Text page"),
        (2, "First line\nSecond line"),
        (3, "Last page"),
    ]
    assert result.blocks[1].anchor.section_path == ("Chapter 1",)
    assert engine.images[0].startswith(b"\x89PNG")


def test_pages_with_nothing_read_stay_flagged(tmp_path):
    path = pdf_with_scan(tmp_path)
    parsed = parse_document(path, "pdf")

    result = run_ocr(path, parsed, FakeEngine([]))

    assert result.pages_needing_ocr == (2,)
    assert [b.text for b in result.blocks] == ["Text page", "Last page"]


def test_document_without_flagged_pages_is_returned_unchanged(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Hello")
    parsed = parse_document(path, "txt")
    engine = FakeEngine()
    assert run_ocr(path, parsed, engine) is parsed
    assert engine.images == []


def test_pptx_picture_slides_keep_their_title_as_section(tmp_path):
    presentation = pptx.Presentation()
    text = presentation.slides.add_slide(presentation.slide_layouts[5])
    text.shapes.title.text = "Intro"
    diagram = presentation.slides.add_slide(presentation.slide_layouts[5])
    diagram.shapes.title.text = "Architecture"
    diagram.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(4))
    diagram.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(2))
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    parsed = parse_document(path, "pptx")
    assert parsed.pages_needing_ocr == (2,)

    result = run_ocr(path, parsed, FakeEngine(["upper picture"], ["lower picture"]))

    slide_two = [b for b in result.blocks if b.anchor.slide == 2]
    assert [b.text for b in slide_two] == ["Architecture", "upper picture", "lower picture"]
    assert slide_two[1].anchor.section_path == ("Architecture",)
    assert result.pages_needing_ocr == ()


def test_ocr_text_is_normalized(tmp_path):
    path = pdf_with_scan(tmp_path)
    parsed = parse_document(path, "pdf")
    result = run_ocr(path, parsed, FakeEngine(["Cafe\u0301\x00 "]))
    assert result.blocks[1].text == "Café"


@pytest.mark.skipif(
    importlib.util.find_spec("paddleocr") is None,
    reason="PaddleOCR not installed (uv sync --extra ocr)",
)
def test_paddle_engine_reads_rendered_text():
    from app.ingestion.ocr import PaddleOcrEngine

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 100), "Hello InsightHub", fontsize=32)
    image = doc[0].get_pixmap(dpi=200).tobytes("png")
    assert "InsightHub" in " ".join(PaddleOcrEngine("en").recognize(image))
