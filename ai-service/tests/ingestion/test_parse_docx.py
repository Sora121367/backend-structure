import docx
import pytest

from app.ingestion.parsing import DocumentParseError, parse_document

KHMER = "សួស្តី\u200bពិភពលោក"


def save(document, tmp_path):
    path = tmp_path / "doc.docx"
    document.save(str(path))
    return path


def test_headings_set_nested_section_paths(tmp_path):
    document = docx.Document()
    document.add_heading("Course notes", 0)  # Title style
    document.add_heading("Chapter 1", 1)
    document.add_paragraph("Intro text")
    document.add_heading("1.1 Basics", 2)
    document.add_paragraph("Basics text")
    document.add_heading("Chapter 2", 1)
    document.add_paragraph("More text")
    parsed = parse_document(save(document, tmp_path), "docx")

    assert parsed.title == "Course notes"
    assert [(b.kind, b.heading_level, b.text, b.anchor.section_path) for b in parsed.blocks] == [
        ("heading", 0, "Course notes", ()),
        ("heading", 1, "Chapter 1", ("Chapter 1",)),
        ("paragraph", None, "Intro text", ("Chapter 1",)),
        ("heading", 2, "1.1 Basics", ("Chapter 1", "1.1 Basics")),
        ("paragraph", None, "Basics text", ("Chapter 1", "1.1 Basics")),
        ("heading", 1, "Chapter 2", ("Chapter 2",)),
        ("paragraph", None, "More text", ("Chapter 2",)),
    ]
    assert [b.anchor.paragraph_index for b in parsed.blocks] == list(range(7))


def test_core_title_wins_over_title_style(tmp_path):
    document = docx.Document()
    document.core_properties.title = "From properties"
    document.add_heading("From body", 0)
    assert parse_document(save(document, tmp_path), "docx").title == "From properties"


def test_list_items_and_khmer(tmp_path):
    document = docx.Document()
    document.add_paragraph("First point", style="List Bullet")
    document.add_paragraph("Second point", style="List Number")
    document.add_paragraph(KHMER)
    parsed = parse_document(save(document, tmp_path), "docx")

    assert [b.kind for b in parsed.blocks] == ["list_item", "list_item", "paragraph"]
    assert parsed.blocks[2].text == KHMER


def test_tables_keep_body_order_and_merged_cells_once(tmp_path):
    document = docx.Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Header"
    table.cell(1, 0).text, table.cell(1, 1).text = "a", "b"
    document.add_paragraph("After")
    parsed = parse_document(save(document, tmp_path), "docx")

    assert [b.kind for b in parsed.blocks] == ["paragraph", "table", "paragraph"]
    assert parsed.blocks[1].text == "Header\na | b"  # empty third row dropped


def test_rejects_empty_document(tmp_path):
    with pytest.raises(DocumentParseError, match="no text"):
        parse_document(save(docx.Document(), tmp_path), "docx")


@pytest.mark.parametrize("data", [b"", b"not a docx", b"PK\x03\x04broken"])
def test_rejects_unreadable_files(tmp_path, data):
    path = tmp_path / "bad.docx"
    path.write_bytes(data)
    with pytest.raises(DocumentParseError):
        parse_document(path, "docx")
